import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin
from pypdf import PdfReader
from io import BytesIO
import re
import time
from openpyxl import Workbook


# ============================================================
# CONFIGURACIÓN
# ============================================================

URL_BASE = (
    "https://www.gob.pe/institucion/oefa/colecciones/"
    "1716-sala-especializada-en-mineria-energia-actividades-productivas"
    "-e-infraestructura-y-servicios-tfa-se"
)

MAX_PAGINAS = 19

encabezados = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/131.0 Safari/537.36"
    )
}


# ============================================================
# FUNCIONES
# ============================================================

TEMAS = {

    # ========================================================
    # COMPONENTES AMBIENTALES
    # ========================================================

    "ruido": [
        "ruido",
        "ruido ambiental",
        "ruido ocupacional",
        "nivel de presión sonora",
        "presión sonora",
        "decibel",
        "decibeles",
        "db(a)",
        "dba",
        "contaminación sonora",
        "emisión sonora"
    ],

    "aire": [
        "calidad del aire",
        "aire ambiente",
        "calidad ambiental del aire",
        "material particulado",
        "pm10",
        "pm2.5",
        "pm 10",
        "pm 2.5",
        "partículas",
        "particulas",
        "polvo",
        "inmisión",
        "inmision"
    ],

    "agua": [
        "calidad del agua",
        "recurso hídrico",
        "recurso hidrico",
        "cuerpo de agua",
        "cuerpos de agua",
        "agua superficial",
        "agua subterránea",
        "agua subterranea",
        "agua residual",
        "agua de mar",
        "agua continental",
        "hidroquímica",
        "hidroquimica"
    ],

    "suelo": [
        "suelo",
        "calidad del suelo",
        "contaminación del suelo",
        "contaminacion del suelo",
        "suelo contaminado",
        "suelo industrial",
        "suelo agrícola",
        "suelo agricola"
    ],

    "sedimentos": [
        "sedimento",
        "sedimentos",
        "sedimento superficial",
        "sedimentos superficiales",
        "calidad de sedimentos",
        "contaminación de sedimentos",
        "contaminacion de sedimentos"
    ],

    "vibraciones": [
        "vibración",
        "vibraciones",
        "vibración ambiental",
        "vibraciones ambientales"
    ],

    "olores": [
        "olor",
        "olores",
        "olor ambiental",
        "emisión de olores",
        "contaminación odorífera",
        "contaminacion odorifera"
    ],

    "flora": [
        "flora",
        "vegetación",
        "vegetacion",
        "cobertura vegetal",
        "especie vegetal",
        "especies vegetales",
        "vegetal"
    ],

    "fauna": [
        "fauna",
        "fauna silvestre",
        "especie de fauna",
        "especies de fauna",
        "especie animal",
        "especies animales"
    ],

    "biodiversidad": [
        "biodiversidad",
        "diversidad biológica",
        "diversidad biologica",
        "ecosistema",
        "ecosistemas",
        "hábitat",
        "habitat",
        "área natural protegida",
        "area natural protegida"
    ],


    # ========================================================
    # FUENTES Y ASPECTOS AMBIENTALES
    # ========================================================

    "emisiones": [
        "emisiones",
        "emisión",
        "emision",
        "emisiones atmosféricas",
        "emisiones atmosfericas",
        "emisión atmosférica",
        "emision atmosferica",
        "fuente de emisión",
        "fuente de emision",
        "gases",
        "gases contaminantes",
        "chimenea",
        "fuente fija",
        "fuente móvil",
        "fuente movil"
    ],

    "efluentes": [
        "efluente",
        "efluentes",
        "efluente líquido",
        "efluente liquido",
        "vertimiento",
        "vertimientos",
        "descarga",
        "descargas",
        "agua residual industrial",
        "aguas residuales industriales",
        "punto de vertimiento"
    ],

    "relaves": [
        "relave",
        "relaves",
        "depósito de relaves",
        "deposito de relaves",
        "depósito de relave",
        "deposito de relave",
        "presa de relaves",
        "relavera",
        "relaveras"
    ],

    "residuos": [
        "residuo",
        "residuos",
        "residuos sólidos",
        "residuos solidos",
        "residuo sólido",
        "residuo solido",
        "residuos peligrosos",
        "residuo peligroso",
        "manejo de residuos",
        "gestión de residuos",
        "gestion de residuos",
        "disposición final",
        "disposicion final",
        "botadero"
    ],

    "hidrocarburos": [
        "hidrocarburo",
        "hidrocarburos",
        "petróleo",
        "petroleo",
        "combustible",
        "combustibles",
        "aceite",
        "aceites",
        "lubricante",
        "lubricantes",
        "derrame de hidrocarburos",
        "derrame de combustible"
    ],


    # ========================================================
    # GESTIÓN AMBIENTAL
    # ========================================================

    "monitoreo ambiental": [
        "monitoreo ambiental",
        "monitoreo",
        "monitoreo ambiental participativo",
        "vigilancia ambiental",
        "medición ambiental",
        "mediciones ambientales",
        "punto de monitoreo",
        "puntos de monitoreo",
        "estación de monitoreo",
        "estaciones de monitoreo"
    ],

    "instrumento de gestión ambiental": [
        "instrumento de gestión ambiental",
        "instrumento de gestion ambiental",
        "iga",
        "estudio de impacto ambiental",
        "eia",
        "declaración de impacto ambiental",
        "declaracion de impacto ambiental",
        "plan de manejo ambiental",
        "pma",
        "programa de adecuación y manejo ambiental",
        "programa de adecuacion y manejo ambiental",
        "plan de cierre",
        "instrumento ambiental"
    ],

    "eca": [
        "estándar de calidad ambiental",
        "estandar de calidad ambiental",
        "estándares de calidad ambiental",
        "estandares de calidad ambiental",
        "eca",
        "eca para agua",
        "eca para aire",
        "eca para suelo",
        "eca ruido"
    ],

    "lmp": [
        "límite máximo permisible",
        "limite maximo permisible",
        "límites máximos permisibles",
        "limites maximos permisibles",
        "lmp",
        "lmp de emisiones",
        "lmp de efluentes"
    ],

    "obligaciones ambientales": [
        "obligación ambiental",
        "obligaciones ambientales",
        "obligación ambiental fiscalizable",
        "obligaciones ambientales fiscalizables",
        "compromiso ambiental",
        "compromisos ambientales",
        "obligación establecida",
        "obligaciones establecidas"
    ],


    # ========================================================
    # FISCALIZACIÓN Y PROCEDIMIENTO
    # ========================================================

    "fiscalización": [
        "fiscalización ambiental",
        "fiscalizacion ambiental",
        "supervisión ambiental",
        "supervision ambiental",
        "evaluación ambiental",
        "evaluacion ambiental",
        "acción de supervisión",
        "accion de supervision",
        "informe de supervisión",
        "informe de supervision"
    ],

    "incumplimiento": [
        "incumplimiento",
        "incumplió",
        "incumplio",
        "incumple",
        "no cumplió",
        "no cumplio",
        "incumplimiento de la obligación",
        "incumplimiento de obligaciones"
    ],

    "infracción": [
        "infracción",
        "infraccion",
        "infracción administrativa",
        "infraccion administrativa",
        "conducta infractora",
        "conducta infractora"
    ],

    "sanción": [
        "sanción",
        "sancion",
        "sanción administrativa",
        "sancion administrativa",
        "multa",
        "multa administrativa",
        "sancionado",
        "sancionada"
    ],

    "medidas administrativas": [
        "medida administrativa",
        "medidas administrativas",
        "medida preventiva",
        "medidas preventivas",
        "medida cautelar",
        "medidas cautelares",
        "mandato de carácter particular",
        "mandato de caracter particular",
        "medida correctiva",
        "medidas correctivas"
    ],

    "beneficio ilícito": [
        "beneficio ilícito",
        "beneficio ilicito",
        "beneficio ilícito obtenido",
        "beneficio ilicito obtenido",
        "cálculo del beneficio ilícito",
        "calculo del beneficio ilicito",
        "beneficio ilícito (b)"
    ],

    "daño ambiental": [
        "daño ambiental",
        "dano ambiental",
        "daño al ambiente",
        "dano al ambiente",
        "afectación ambiental",
        "afectacion ambiental",
        "afectación al ambiente",
        "afectacion al ambiente"
    ]
}

def normalizar(texto):
    texto = texto.lower()
    texto = texto.replace("á", "a")
    texto = texto.replace("é", "e")
    texto = texto.replace("í", "i")
    texto = texto.replace("ó", "o")
    texto = texto.replace("ú", "u")
    texto = texto.replace("ü", "u")
    return texto


def obtener_resoluciones():
    """
    Recorre las páginas de la colección del TFA
    y obtiene las URL de las resoluciones.
    """

    resoluciones = []

    print()
    print("=" * 70)
    print("OBTENIENDO RESOLUCIONES DEL TFA-OEFA")
    print("=" * 70)

    for numero_pagina in range(1, MAX_PAGINAS + 1):

        print()
        print(f"Leyendo página {numero_pagina} de {MAX_PAGINAS}...")

        if numero_pagina == 1:
            url = URL_BASE
        else:
            url = f"{URL_BASE}?sheet={numero_pagina}"

        try:
            respuesta = requests.get(
                url,
                headers=encabezados,
                timeout=30
            )

            print("Código de respuesta:", respuesta.status_code)

            if respuesta.status_code != 200:
                continue

            sopa = BeautifulSoup(respuesta.text, "html.parser")

            enlaces = sopa.find_all("a", href=True)

            cantidad_antes = len(resoluciones)

            for enlace in enlaces:

                texto = enlace.get_text(" ", strip=True)
                href = enlace["href"]

                url_completa = urljoin(url, href)

                # Solo nos interesan páginas de informes-publicaciones
                if "/informes-publicaciones/" in url_completa:

                    # Verificamos que corresponda a una resolución TFA
                    if "resolucion" in normalizar(texto):

                        if url_completa not in resoluciones:
                            resoluciones.append(url_completa)

            nuevas = len(resoluciones) - cantidad_antes

            print("Nuevas resoluciones encontradas:", nuevas)
            print("Total acumulado:", len(resoluciones))

            time.sleep(0.5)

        except Exception as error:
            print("Error:", error)

    return resoluciones


def obtener_pdf(url_resolucion):
    """
    Entra a la página de la resolución y obtiene
    la URL del PDF.
    """

    try:

        respuesta = requests.get(
            url_resolucion,
            headers=encabezados,
            timeout=30
        )

        if respuesta.status_code != 200:
            return None

        sopa = BeautifulSoup(respuesta.text, "html.parser")

        for enlace in sopa.find_all("a", href=True):

            href = enlace["href"]

            if ".pdf" in href.lower():

                return urljoin(url_resolucion, href)

        return None

    except Exception:
        return None


# ============================================================
# V4 - ANALIZADOR DOCUMENTAL TFA-OEFA
# ============================================================
#
# Herramienta para búsqueda y análisis documental de resoluciones
# del Tribunal de Fiscalización Ambiental (TFA-OEFA).
#
# Características principales:
# - Procesamiento de PDF únicamente en memoria mediante BytesIO.
# - No guarda ni descarga PDFs en disco.
# - Búsqueda por temas, varios temas, todos los temas o texto directo.
# - Identificación del administrado a partir del encabezado del PDF.
# - Evidencia textual conservada desde la extracción original del PDF.
# - Resumen, detalle, coocurrencia, administrados y gráficos en Excel.
# ============================================================

import re
import time
import unicodedata
from io import BytesIO
from itertools import combinations

import requests
from pypdf import PdfReader
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule


# ============================================================
# CONFIGURACIÓN DE SALIDAS
# ============================================================

VERSION = "4.0"
ARCHIVO_RESULTADOS = "resultados_busqueda_v4.xlsx"
ARCHIVO_ANALISIS = "analisis_temas_v4.xlsx"


# ============================================================
# NORMALIZACIÓN PARA BÚSQUEDA
# ============================================================

def normalizar_busqueda(texto):
    """
    Normaliza solamente para realizar la búsqueda.
    El texto original del PDF nunca se modifica para la evidencia.
    """
    texto = texto.lower()
    texto = unicodedata.normalize("NFD", texto)
    texto = "".join(
        caracter
        for caracter in texto
        if unicodedata.category(caracter) != "Mn"
    )
    return texto


def normalizar_con_mapeo(texto_original):
    """
    Crea un texto normalizado y un mapa que permite volver desde
    una posición normalizada hasta la posición equivalente del texto original.
    """
    texto_normalizado = []
    mapa = []

    for posicion_original, caracter in enumerate(texto_original):
        caracter_normalizado = unicodedata.normalize(
            "NFD",
            caracter.lower()
        )

        caracter_normalizado = "".join(
            c
            for c in caracter_normalizado
            if unicodedata.category(c) != "Mn"
        )

        for caracter_resultante in caracter_normalizado:
            texto_normalizado.append(caracter_resultante)
            mapa.append(posicion_original)

    return "".join(texto_normalizado), mapa


# ============================================================
# PATRONES DE BÚSQUEDA
# ============================================================

def crear_patron(termino):
    """Crea una expresión regular para palabras o frases completas."""
    termino_normalizado = normalizar_busqueda(termino).strip()

    if not termino_normalizado:
        return None

    partes = termino_normalizado.split()

    if len(partes) == 1:
        expresion = (
            r"(?<!\w)"
            + re.escape(partes[0])
            + r"(?!\w)"
        )
    else:
        expresion = (
            r"(?<!\w)"
            + r"\s+".join(re.escape(parte) for parte in partes)
            + r"(?!\w)"
        )

    return re.compile(expresion)


def preparar_patrones(terminos_por_tema):
    """
    Compila los patrones una sola vez para acelerar el procesamiento.
    """
    patrones = []

    for tema, termino in terminos_por_tema:
        patron = crear_patron(termino)
        if patron is not None:
            patrones.append((tema, termino, patron))

    return patrones


def seleccionar_coincidencias_no_superpuestas(texto_normalizado, patrones):
    """
    Busca todos los términos de un tema y evita contar como coincidencias
    independientes aquellas que se superponen.

    Se prioriza la coincidencia más larga. Ejemplo:
    'calidad del agua' tiene prioridad sobre 'agua' en el mismo lugar.
    """
    candidatas = []

    for termino, patron in patrones:
        for coincidencia in patron.finditer(texto_normalizado):
            candidatas.append(
                {
                    "termino": termino,
                    "inicio": coincidencia.start(),
                    "fin": coincidencia.end(),
                    "longitud": coincidencia.end() - coincidencia.start(),
                }
            )

    candidatas.sort(
        key=lambda x: (
            -x["longitud"],
            x["inicio"],
            x["termino"]
        )
    )

    seleccionadas = []

    for candidata in candidatas:
        se_superpone = any(
            candidata["inicio"] < existente["fin"]
            and candidata["fin"] > existente["inicio"]
            for existente in seleccionadas
        )

        if not se_superpone:
            seleccionadas.append(candidata)

    seleccionadas.sort(
        key=lambda x: (x["inicio"], x["fin"])
    )

    return seleccionadas


# ============================================================
# CONTEXTO ORIGINAL
# ============================================================

def extraer_contexto(texto_original, inicio, fin):
    """
    Extrae contexto desde el texto original del PDF.
    No cambia mayúsculas, tildes ni redacción.
    """
    caracteres_antes = 180
    caracteres_despues = 250

    inicio_contexto = max(
        0,
        inicio - caracteres_antes
    )

    fin_contexto = min(
        len(texto_original),
        fin + caracteres_despues
    )

    if inicio_contexto > 0:
        espacio = texto_original.find(" ", inicio_contexto)
        if espacio != -1 and espacio < inicio:
            inicio_contexto = espacio + 1

    if fin_contexto < len(texto_original):
        espacio = texto_original.rfind(
            " ",
            inicio,
            fin_contexto
        )
        if espacio != -1:
            fin_contexto = espacio

    return texto_original[inicio_contexto:fin_contexto].strip()


# ============================================================
# SECCIÓN APROXIMADA
# ============================================================

def detectar_seccion(texto_original, posicion):
    """
    Identifica de manera aproximada un encabezado textual cercano.
    No constituye una interpretación jurídica.
    """
    texto_anterior = texto_original[:posicion]
    lineas = texto_anterior.splitlines()

    posibles_secciones = [
        "antecedentes",
        "cuestiones controvertidas",
        "cuestión controvertida",
        "análisis de las cuestiones controvertidas",
        "análisis",
        "fundamentos",
        "determinación de responsabilidad",
        "determinación de la responsabilidad",
        "graduación de la multa",
        "multa",
        "medidas administrativas",
        "parte resolutiva",
        "resuelve",
        "decisión",
    ]

    for linea in reversed(lineas[-40:]):
        linea_limpia = normalizar_busqueda(linea).strip()

        if not linea_limpia or len(linea_limpia) > 120:
            continue

        for seccion in posibles_secciones:
            if seccion in linea_limpia:
                return linea.strip()

    return "No identificada"


# ============================================================
# IDENTIFICACIÓN DEL ADMINISTRADO
# ============================================================

ENCABEZADOS_ADMINISTRADO = {
    "expediente",
    "procedencia",
    "administrado",
    "administrados",
    "sector",
    "materia",
    "sumilla",
    "resolucion",
    "resolución",
    "fecha",
    "asunto",
}

def limpiar_texto_extraido(texto):
    """Reduce espacios y saltos de línea sin cambiar las palabras."""
    return re.sub(r"\s+", " ", texto).strip()


def es_encabezado_ficha(linea):
    """Indica si una línea parece ser otro campo del encabezado de la resolución."""
    limpia = normalizar_busqueda(linea).strip().rstrip(":")
    if not limpia:
        return False

    if limpia in {normalizar_busqueda(x) for x in ENCABEZADOS_ADMINISTRADO}:
        return True

    return bool(
        re.match(
            r"^(expediente|procedencia|administrado|administrados|sector|"
            r"materia|sumilla|resolucion|fecha|asunto)\s*:",
            limpia,
        )
    )


def extraer_administrado(texto_original):
    """
    Identifica el campo ADMINISTRADO/ADMINISTRADOS del encabezado del PDF.

    La función busca primero una etiqueta explícita y, si el nombre está
    dividido en varias líneas, concatena las líneas siguientes hasta encontrar
    otro campo del encabezado.

    El valor devuelto conserva el texto extraído del PDF; únicamente se
    normalizan espacios y saltos de línea para facilitar su lectura.
    """
    if not texto_original or not texto_original.strip():
        return "No identificado"

    lineas = [
        linea.strip()
        for linea in texto_original.replace("\r", "\n").splitlines()
    ]

    for indice, linea in enumerate(lineas):
        if not linea:
            continue

        linea_normalizada = normalizar_busqueda(linea).strip()

        # Caso habitual: ADMINISTRADO: NOMBRE
        coincidencia = re.match(
            r"^administrados?\s*:\s*(.*)$",
            linea_normalizada,
        )

        if coincidencia:
            # Recuperar el contenido desde la línea original, no desde
            # la versión normalizada, para conservar tildes y mayúsculas.
            contenido = re.sub(
                r"^\s*administrados?\s*:\s*",
                "",
                linea,
                flags=re.IGNORECASE,
            ).strip()

            partes = [contenido] if contenido else []

            # Si el nombre continúa en las líneas siguientes, incorporarlo.
            for siguiente in lineas[indice + 1:indice + 5]:
                if not siguiente:
                    continue
                if es_encabezado_ficha(siguiente):
                    break
                partes.append(siguiente)

            resultado = limpiar_texto_extraido(" ".join(partes))
            if resultado:
                return resultado

        # Caso: ADMINISTRADO sin dos puntos, seguido del nombre.
        if linea_normalizada.rstrip(":") in {
            "administrado",
            "administrados",
        }:
            partes = []

            for siguiente in lineas[indice + 1:indice + 5]:
                if not siguiente:
                    continue
                if es_encabezado_ficha(siguiente):
                    break
                partes.append(siguiente)

            resultado = limpiar_texto_extraido(" ".join(partes))
            if resultado:
                return resultado

    return "No identificado"


# ============================================================
# ANALIZAR PDF COMPLETO EN MEMORIA
# ============================================================

def analizar_pdf_completo(url_pdf, patrones_por_tema):
    """
    Descarga el PDF únicamente en memoria mediante BytesIO.
    Nunca guarda el PDF en disco.

    Además de las coincidencias temáticas, identifica el administrado
    a partir del encabezado de las primeras páginas del documento.
    """
    archivo = None

    try:
        print("    Descargando PDF en memoria...")

        respuesta = requests.get(
            url_pdf,
            headers=encabezados,
            timeout=60
        )
        respuesta.raise_for_status()

        archivo = BytesIO(respuesta.content)
        lector = PdfReader(archivo)
        total_paginas = len(lector.pages)
        coincidencias = []
        paginas_con_texto = 0

        print(f"    Páginas del PDF: {total_paginas}")

        # El encabezado con el administrado suele estar en la primera página.
        # Se revisan las primeras 3 páginas para tolerar PDFs con carátulas.
        textos_iniciales = []

        # Procesar tema por tema para controlar las superposiciones.
        patrones_por_tema_agrupados = {}
        for tema, termino, patron in patrones_por_tema:
            patrones_por_tema_agrupados.setdefault(
                tema,
                []
            ).append((termino, patron))

        for numero_pagina, pagina in enumerate(lector.pages, 1):
            try:
                texto_original = pagina.extract_text() or ""
            except Exception:
                texto_original = ""

            if numero_pagina <= 3 and texto_original.strip():
                textos_iniciales.append(texto_original)

            if not texto_original.strip():
                continue

            paginas_con_texto += 1

            texto_normalizado, mapa_posiciones = normalizar_con_mapeo(
                texto_original
            )

            for tema, patrones in patrones_por_tema_agrupados.items():
                coincidencias_tema = seleccionar_coincidencias_no_superpuestas(
                    texto_normalizado,
                    patrones
                )

                for coincidencia in coincidencias_tema:
                    inicio_normalizado = coincidencia["inicio"]
                    fin_normalizado = coincidencia["fin"]

                    if inicio_normalizado >= len(mapa_posiciones):
                        continue

                    inicio_original = mapa_posiciones[inicio_normalizado]

                    if (
                        fin_normalizado - 1 < len(mapa_posiciones)
                    ):
                        fin_original = (
                            mapa_posiciones[fin_normalizado - 1] + 1
                        )
                    else:
                        fin_original = len(texto_original)

                    contexto = extraer_contexto(
                        texto_original,
                        inicio_original,
                        fin_original
                    )

                    seccion = detectar_seccion(
                        texto_original,
                        inicio_original
                    )

                    coincidencias.append(
                        {
                            "tema": tema,
                            "termino": coincidencia["termino"],
                            "pagina": numero_pagina,
                            "contexto": contexto,
                            "seccion": seccion,
                        }
                    )

        administrado = extraer_administrado(
            "\n".join(textos_iniciales)
        )

        return {
            "estado": "Procesado",
            "error": "",
            "administrado": administrado,
            "total_paginas": total_paginas,
            "paginas_con_texto": paginas_con_texto,
            "coincidencias": coincidencias,
        }

    except Exception as error:
        print(f"    ERROR leyendo PDF: {error}")
        return {
            "estado": "Error de lectura",
            "error": str(error),
            "administrado": "No identificado",
            "total_paginas": 0,
            "paginas_con_texto": 0,
            "coincidencias": [],
        }

    finally:
        if archivo is not None:
            archivo.close()


# ============================================================
# SELECCIONAR TEMAS Y TÉRMINOS
# ============================================================

def preparar_terminos_busqueda(entrada):
    """Permite tema, varios temas, todos o búsqueda directa."""
    entrada = entrada.strip()

    if not entrada:
        return []

    partes = [
        parte.strip()
        for parte in entrada.split(",")
        if parte.strip()
    ]

    terminos_por_tema = []

    if any(
        normalizar_busqueda(parte) == "todos"
        for parte in partes
    ):
        for tema, terminos in TEMAS.items():
            for termino in terminos:
                item = (tema, termino)
                if item not in terminos_por_tema:
                    terminos_por_tema.append(item)
        return terminos_por_tema

    for parte in partes:
        parte_normalizada = normalizar_busqueda(parte).strip()
        tema_encontrado = None

        for tema in TEMAS:
            if normalizar_busqueda(tema).strip() == parte_normalizada:
                tema_encontrado = tema
                break

        if tema_encontrado is not None:
            for termino in TEMAS[tema_encontrado]:
                item = (tema_encontrado, termino)
                if item not in terminos_por_tema:
                    terminos_por_tema.append(item)
        else:
            item = ("Búsqueda directa", parte)
            if item not in terminos_por_tema:
                terminos_por_tema.append(item)

    return terminos_por_tema


def obtener_anio(nombre):
    coincidencia = re.search(r"20\d{2}", nombre)
    return int(coincidencia.group()) if coincidencia else ""


# ============================================================
# INICIO DEL PROGRAMA
# ============================================================

inicio_proceso = time.time()

print()
print("=" * 80)
print(f"ANALIZADOR V{VERSION} DE RESOLUCIONES TFA-OEFA")
print("=" * 80)
print()
print("Puedes buscar:")
print("  - un tema: ruido")
print("  - varios temas: ruido, agua, emisiones")
print("  - todos los temas: todos")
print("  - una palabra o frase directa")
print()

entrada = input(
    "Escribe el tema o temas que quieres analizar: "
)

if entrada.strip() == "":
    print("No escribiste ningún término.")
    exit()

terminos_busqueda = preparar_terminos_busqueda(entrada)

if not terminos_busqueda:
    print("No se encontraron términos para analizar.")
    exit()

patrones = preparar_patrones(terminos_busqueda)

print()
print("=" * 80)
print("TÉRMINOS QUE SERÁN ANALIZADOS")
print("=" * 80)

for tema, termino in terminos_busqueda:
    print(f"  - {tema}: {termino}")

print()
print("Patrones preparados:", len(patrones))


# ============================================================
# OBTENER RESOLUCIONES
# ============================================================

print("=" * 80)
print("OBTENIENDO RESOLUCIONES")
print("=" * 80)

resoluciones = obtener_resoluciones()

print()
print("Total de resoluciones encontradas:", len(resoluciones))

if not resoluciones:
    print("No se encontraron resoluciones.")
    exit()


# ============================================================
# VARIABLES GENERALES
# ============================================================

resultados_detalle = []
resumen_resoluciones = []
temas_resumen = {}

for tema, termino in terminos_busqueda:
    if tema not in temas_resumen:
        temas_resumen[tema] = {
            "resoluciones": set(),
            "coincidencias": 0,
            "paginas": set(),
        }

resoluciones_con_coincidencias = 0
total_coincidencias = 0
total_paginas_con_coincidencias = set()
total_paginas_analizadas = 0
pdf_sin_texto = 0
pdf_sin_coincidencias = 0
pdf_no_encontrado = 0
pdf_error = 0

# Administrados identificados. La clave normalizada permite agrupar
# diferencias menores de mayúsculas/minúsculas sin alterar el nombre
# mostrado en Excel.
administrados_resumen = {}
resoluciones_sin_administrado = 0


# ============================================================
# PROCESAR CADA RESOLUCIÓN
# ============================================================

print()
print("=" * 80)
print("INICIANDO ANÁLISIS")
print("=" * 80)
print()

for numero, url_resolucion in enumerate(resoluciones, 1):
    print()
    print(f"[{numero}/{len(resoluciones)}]")

    nombre = (
        url_resolucion
        .rstrip("/")
        .split("/")[-1]
    )

    print("Resolución:", nombre)
    anio = obtener_anio(nombre)
    url_pdf = obtener_pdf(url_resolucion)

    if not url_pdf:
        print("    No se encontró PDF.")
        pdf_no_encontrado += 1

        resumen_resoluciones.append(
            {
                "Resolución": nombre,
                "Año": anio,
                "URL Resolución": url_resolucion,
                "URL PDF": "",
                "Administrado": "No identificado",
                "PDF": "No encontrado",
                "Páginas PDF": 0,
                "Páginas con texto": 0,
                "Páginas con coincidencias": 0,
                "Total coincidencias": 0,
                **{tema: 0 for tema in temas_resumen},
            }
        )
        continue

    print("    PDF encontrado.")

    resultado_pdf = analizar_pdf_completo(
        url_pdf,
        patrones
    )

    estado_pdf = resultado_pdf["estado"]
    administrado = resultado_pdf.get("administrado", "No identificado")
    coincidencias = resultado_pdf["coincidencias"]
    total_paginas = resultado_pdf["total_paginas"]
    paginas_con_texto = resultado_pdf["paginas_con_texto"]

    if administrado == "No identificado":
        resoluciones_sin_administrado += 1
    else:
        clave_administrado = normalizar_busqueda(administrado).strip()
        datos_administrado = administrados_resumen.setdefault(
            clave_administrado,
            {
                "Administrado": administrado,
                "Resoluciones": set(),
                "Coincidencias": 0,
                "Años": set(),
            }
        )
        datos_administrado["Resoluciones"].add(nombre)
        if anio:
            datos_administrado["Años"].add(anio)

    total_paginas_analizadas += total_paginas

    if estado_pdf == "Error de lectura":
        pdf_error += 1
    elif paginas_con_texto == 0:
        pdf_sin_texto += 1

    paginas_coincidencias = {
        coincidencia["pagina"]
        for coincidencia in coincidencias
    }

    if coincidencias:
        resoluciones_con_coincidencias += 1
    elif estado_pdf == "Procesado":
        pdf_sin_coincidencias += 1

    total_coincidencias += len(coincidencias)

    total_paginas_con_coincidencias.update(
        (nombre, pagina)
        for pagina in paginas_coincidencias
    )

    conteo_temas_resolucion = {
        tema: 0
        for tema in temas_resumen
    }

    for coincidencia in coincidencias:
        tema = coincidencia["tema"]
        conteo_temas_resolucion[tema] = (
            conteo_temas_resolucion.get(tema, 0) + 1
        )

        temas_resumen[tema]["resoluciones"].add(nombre)
        temas_resumen[tema]["coincidencias"] += 1
        temas_resumen[tema]["paginas"].add(
            (nombre, coincidencia["pagina"])
        )

        if administrado != "No identificado":
            clave_administrado = normalizar_busqueda(administrado).strip()
            if clave_administrado in administrados_resumen:
                administrados_resumen[clave_administrado]["Coincidencias"] += 1

        resultados_detalle.append(
            {
                "Resolución": nombre,
                "Año": anio,
                "URL Resolución": url_resolucion,
                "URL PDF": url_pdf,
                "Administrado": administrado,
                "Tema": coincidencia["tema"],
                "Término encontrado": coincidencia["termino"],
                "Página": coincidencia["pagina"],
                "Sección aproximada": coincidencia["seccion"],
                "Fragmento": coincidencia["contexto"],
            }
        )

    resumen_resoluciones.append(
        {
            "Resolución": nombre,
            "Año": anio,
            "URL Resolución": url_resolucion,
            "URL PDF": url_pdf,
            "Administrado": administrado,
            "PDF": estado_pdf,
            "Páginas PDF": total_paginas,
            "Páginas con texto": paginas_con_texto,
            "Páginas con coincidencias": len(paginas_coincidencias),
            "Total coincidencias": len(coincidencias),
            **{
                tema: conteo_temas_resolucion.get(tema, 0)
                for tema in temas_resumen
            },
        }
    )

    print(f"    Coincidencias efectivas: {len(coincidencias)}")
    time.sleep(0.2)


# ============================================================
# RESUMEN EN CONSOLA
# ============================================================

print()
print("=" * 80)
print("ANÁLISIS FINALIZADO")
print("=" * 80)
print()
print("Resoluciones encontradas:", len(resoluciones))
print(
    "Resoluciones procesadas:",
    sum(1 for r in resumen_resoluciones if r["PDF"] == "Procesado")
)
print("PDF no encontrados:", pdf_no_encontrado)
print("PDF con error de lectura:", pdf_error)
print("Resoluciones con coincidencias:", resoluciones_con_coincidencias)
print("Resoluciones sin coincidencias:", pdf_sin_coincidencias)
print("PDF sin texto extraíble:", pdf_sin_texto)
print("Resoluciones con administrado identificado:",
      len(resoluciones) - resoluciones_sin_administrado - pdf_no_encontrado)
print("Resoluciones sin administrado identificado:", resoluciones_sin_administrado)
print("Administrados únicos identificados:", len(administrados_resumen))
print("Total de páginas de PDF procesadas:", total_paginas_analizadas)
print("Total de coincidencias efectivas:", total_coincidencias)
print("Páginas con coincidencias:", len(total_paginas_con_coincidencias))
print()
print("RESUMEN POR TEMA:")

for tema, datos in temas_resumen.items():
    print(
        f"  {tema}: "
        f"{len(datos['resoluciones'])} resoluciones | "
        f"{datos['coincidencias']} coincidencias | "
        f"{len(datos['paginas'])} páginas"
    )


# ============================================================
# EXCEL 1: RESULTADOS DE BÚSQUEDA
# ============================================================

print()
print("=" * 80)
print("GENERANDO RESULTADOS DE BÚSQUEDA")
print("=" * 80)

wb_resultados = Workbook()
ws_resultados = wb_resultados.active
ws_resultados.title = "Resultados"

encabezados_resultados = [
    "Resolución",
    "Año",
    "URL Resolución",
    "URL PDF",
    "Administrado",
    "Tema",
    "Término encontrado",
    "Página",
    "Sección aproximada",
    "Fragmento",
]

ws_resultados.append(encabezados_resultados)

for fila in resultados_detalle:
    ws_resultados.append([
        fila["Resolución"],
        fila["Año"],
        fila["URL Resolución"],
        fila["URL PDF"],
        fila["Administrado"],
        fila["Tema"],
        fila["Término encontrado"],
        fila["Página"],
        fila["Sección aproximada"],
        fila["Fragmento"],
    ])

for celda in ws_resultados[1]:
    celda.font = Font(bold=True)
    celda.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

anchos_resultados = {
    "A": 45,
    "B": 10,
    "C": 60,
    "D": 80,
    "E": 45,
    "F": 28,
    "G": 35,
    "H": 12,
    "I": 35,
    "J": 110,
}

for columna, ancho in anchos_resultados.items():
    ws_resultados.column_dimensions[columna].width = ancho

for fila in ws_resultados.iter_rows(
    min_row=2,
    max_row=ws_resultados.max_row
):
    fila[9].alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )

    fila[2].hyperlink = fila[2].value
    fila[3].hyperlink = fila[3].value

ws_resultados.freeze_panes = "A2"
ws_resultados.auto_filter.ref = ws_resultados.dimensions

wb_resultados.save(ARCHIVO_RESULTADOS)
print(f"Archivo generado: {ARCHIVO_RESULTADOS}")


# ============================================================
# EXCEL 2: ANÁLISIS TEMÁTICO
# ============================================================

print()
print("=" * 80)
print("GENERANDO ANÁLISIS TEMÁTICO")
print("=" * 80)

wb = Workbook()


# ============================================================
# HOJA RESUMEN
# ============================================================

ws_resumen = wb.active
ws_resumen.title = "Resumen"

ws_resumen.append(["Indicador", "Valor"])

indicadores = [
    ("Búsqueda realizada", entrada),
    ("Términos/patrones analizados", len(patrones)),
    ("Resoluciones encontradas", len(resoluciones)),
    (
        "Resoluciones procesadas",
        sum(1 for r in resumen_resoluciones if r["PDF"] == "Procesado")
    ),
    ("PDF no encontrados", pdf_no_encontrado),
    ("PDF con error de lectura", pdf_error),
    ("PDF sin texto extraíble", pdf_sin_texto),
    ("Resoluciones con coincidencias", resoluciones_con_coincidencias),
    ("Resoluciones sin coincidencias", pdf_sin_coincidencias),
    ("Resoluciones con administrado identificado",
     len(resoluciones) - resoluciones_sin_administrado - pdf_no_encontrado),
    ("Resoluciones sin administrado identificado", resoluciones_sin_administrado),
    ("Administrados únicos identificados", len(administrados_resumen)),
    ("Total de páginas de PDF procesadas", total_paginas_analizadas),
    ("Páginas con coincidencias", len(total_paginas_con_coincidencias)),
    ("Total de coincidencias efectivas", total_coincidencias),
]

for indicador, valor in indicadores:
    ws_resumen.append([indicador, valor])

ws_resumen.append([])
ws_resumen.append([
    "TEMA",
    "RESOLUCIONES",
    "COINCIDENCIAS",
    "PÁGINAS",
])

fila_temas_inicio = ws_resumen.max_row

for tema, datos in temas_resumen.items():
    ws_resumen.append([
        tema.capitalize(),
        len(datos["resoluciones"]),
        datos["coincidencias"],
        len(datos["paginas"]),
    ])

fila_temas_fin = ws_resumen.max_row

ws_resumen.append([])
ws_resumen.append([
    "Nota metodológica",
    "Las coincidencias efectivas son ocurrencias textuales de los términos definidos después de controlar superposiciones dentro de cada tema. No representan por sí mismas relevancia jurídica, gravedad ni importancia del asunto."
])

ws_resumen.append([
    "Evidencia",
    "Los fragmentos se extraen del texto original obtenido del PDF. La normalización se utiliza solamente para realizar la búsqueda."
])

ws_resumen.append([
    "Almacenamiento",
    "Los PDF se procesan únicamente en memoria mediante BytesIO y no se guardan en disco."
])

for celda in ws_resumen[1]:
    celda.font = Font(bold=True)

for columna, ancho in {
    "A": 45,
    "B": 80,
    "C": 20,
    "D": 20,
}.items():
    ws_resumen.column_dimensions[columna].width = ancho

for fila in ws_resumen.iter_rows():
    for celda in fila:
        celda.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )


# ============================================================
# HOJA RESOLUCIONES
# ============================================================

ws_resoluciones = wb.create_sheet("Resoluciones")

encabezados_resoluciones = [
    "Resolución",
    "Año",
    "URL Resolución",
    "URL PDF",
    "Administrado",
    "PDF",
    "Páginas PDF",
    "Páginas con texto",
    "Páginas con coincidencias",
    "Total coincidencias",
]
encabezados_resoluciones.extend(temas_resumen.keys())
ws_resoluciones.append(encabezados_resoluciones)

for resultado in resumen_resoluciones:
    fila = [
        resultado["Resolución"],
        resultado["Año"],
        resultado["URL Resolución"],
        resultado["URL PDF"],
        resultado["Administrado"],
        resultado["PDF"],
        resultado["Páginas PDF"],
        resultado["Páginas con texto"],
        resultado["Páginas con coincidencias"],
        resultado["Total coincidencias"],
    ]

    for tema in temas_resumen:
        fila.append(resultado.get(tema, 0))

    ws_resoluciones.append(fila)

for celda in ws_resoluciones[1]:
    celda.font = Font(bold=True)
    celda.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

for fila in ws_resoluciones.iter_rows(min_row=2):
    fila[2].hyperlink = fila[2].value
    fila[3].hyperlink = fila[3].value

ws_resoluciones.freeze_panes = "A2"
ws_resoluciones.auto_filter.ref = ws_resoluciones.dimensions

for columna, ancho in {
    "A": 45,
    "B": 10,
    "C": 60,
    "D": 80,
    "E": 45,
    "F": 20,
    "G": 15,
    "H": 18,
    "I": 25,
    "J": 20,
}.items():
    ws_resoluciones.column_dimensions[columna].width = ancho

for columna in range(10, ws_resoluciones.max_column + 1):
    letra = ws_resoluciones.cell(
        row=1,
        column=columna
    ).column_letter
    ws_resoluciones.column_dimensions[letra].width = 18


# ============================================================
# HOJA DETALLE
# ============================================================

ws_detalle = wb.create_sheet("Detalle")

encabezados_detalle = [
    "Resolución",
    "Año",
    "URL Resolución",
    "URL PDF",
    "Administrado",
    "Tema",
    "Término encontrado",
    "Página",
    "Sección aproximada",
    "Fragmento",
]
ws_detalle.append(encabezados_detalle)

for fila in resultados_detalle:
    ws_detalle.append([
        fila["Resolución"],
        fila["Año"],
        fila["URL Resolución"],
        fila["URL PDF"],
        fila["Administrado"],
        fila["Tema"],
        fila["Término encontrado"],
        fila["Página"],
        fila["Sección aproximada"],
        fila["Fragmento"],
    ])

for celda in ws_detalle[1]:
    celda.font = Font(bold=True)
    celda.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

for columna, ancho in {
    "A": 45,
    "B": 10,
    "C": 60,
    "D": 80,
    "E": 45,
    "F": 28,
    "G": 35,
    "H": 12,
    "I": 35,
    "J": 110,
}.items():
    ws_detalle.column_dimensions[columna].width = ancho

for fila in ws_detalle.iter_rows(min_row=2):
    fila[2].hyperlink = fila[2].value
    fila[3].hyperlink = fila[3].value
    fila[9].alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )

ws_detalle.freeze_panes = "A2"
ws_detalle.auto_filter.ref = ws_detalle.dimensions


# ============================================================
# HOJA ADMINISTRADOS
# ============================================================

ws_administrados = wb.create_sheet("Administrados")

ws_administrados.append([
    "Administrado",
    "Resoluciones",
    "Coincidencias",
    "Años",
])

for datos in sorted(
    administrados_resumen.values(),
    key=lambda item: (
        -len(item["Resoluciones"]),
        normalizar_busqueda(item["Administrado"])
    )
):
    ws_administrados.append([
        datos["Administrado"],
        len(datos["Resoluciones"]),
        datos["Coincidencias"],
        ", ".join(str(anio) for anio in sorted(datos["Años"])),
    ])

for celda in ws_administrados[1]:
    celda.font = Font(bold=True)
    celda.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

for columna, ancho in {
    "A": 65,
    "B": 18,
    "C": 18,
    "D": 25,
}.items():
    ws_administrados.column_dimensions[columna].width = ancho

ws_administrados.freeze_panes = "A2"
ws_administrados.auto_filter.ref = ws_administrados.dimensions

# ============================================================
# HOJA COOCCURRENCIA
# ============================================================

ws_cooc = wb.create_sheet("Coocurrencia")
ws_cooc.append([
    "Tema 1",
    "Tema 2",
    "Resoluciones donde aparecen ambos",
])

temas_por_resolucion = {}

for fila in resultados_detalle:
    resolucion = fila["Resolución"]
    tema = fila["Tema"]
    temas_por_resolucion.setdefault(resolucion, set()).add(tema)

conteo_coocurrencia = {}

for temas in temas_por_resolucion.values():
    temas_validos = sorted(temas)

    for tema1, tema2 in combinations(temas_validos, 2):
        clave = (tema1, tema2)
        conteo_coocurrencia[clave] = (
            conteo_coocurrencia.get(clave, 0) + 1
        )

for (tema1, tema2), cantidad in sorted(
    conteo_coocurrencia.items(),
    key=lambda item: (-item[1], item[0][0], item[0][1])
):
    ws_cooc.append([tema1, tema2, cantidad])

for celda in ws_cooc[1]:
    celda.font = Font(bold=True)
    celda.alignment = Alignment(horizontal="center")

for columna, ancho in {
    "A": 32,
    "B": 32,
    "C": 38,
}.items():
    ws_cooc.column_dimensions[columna].width = ancho

ws_cooc.freeze_panes = "A2"
ws_cooc.auto_filter.ref = ws_cooc.dimensions


# ============================================================
# HOJA GRÁFICOS
# ============================================================

ws_graficos = wb.create_sheet("Gráficos")

ws_graficos["A1"] = "Análisis gráfico de las resoluciones TFA-OEFA"
ws_graficos["A1"].font = Font(bold=True, size=16)

ws_graficos["A2"] = (
    "Los gráficos muestran frecuencia documental de los temas "
    "y distribución de resoluciones por administrado."
)
ws_graficos["A2"].alignment = Alignment(
    wrap_text=True
)

# ============================================================
# GRÁFICO 1: RESOLUCIONES POR TEMA
# ============================================================

grafico_resoluciones = BarChart()

grafico_resoluciones.type = "bar"
grafico_resoluciones.style = 10
grafico_resoluciones.title = "Resoluciones por tema"
grafico_resoluciones.y_axis.title = "Tema"
grafico_resoluciones.x_axis.title = "Número de resoluciones"

datos_resoluciones = Reference(
    ws_resumen,
    min_col=2,
    min_row=fila_temas_inicio,
    max_row=fila_temas_fin
)

categorias_temas = Reference(
    ws_resumen,
    min_col=1,
    min_row=fila_temas_inicio + 1,
    max_row=fila_temas_fin
)

grafico_resoluciones.add_data(
    datos_resoluciones,
    titles_from_data=True
)

grafico_resoluciones.set_categories(
    categorias_temas
)

grafico_resoluciones.height = 10
grafico_resoluciones.width = 17
grafico_resoluciones.legend = None

ws_graficos.add_chart(
    grafico_resoluciones,
    "A4"
)


# ============================================================
# GRÁFICO 2: COINCIDENCIAS POR TEMA
# ============================================================

grafico_coincidencias = BarChart()

grafico_coincidencias.type = "bar"
grafico_coincidencias.style = 10
grafico_coincidencias.title = "Coincidencias textuales por tema"
grafico_coincidencias.y_axis.title = "Tema"
grafico_coincidencias.x_axis.title = "Número de coincidencias"

datos_coincidencias = Reference(
    ws_resumen,
    min_col=3,
    min_row=fila_temas_inicio,
    max_row=fila_temas_fin
)

grafico_coincidencias.add_data(
    datos_coincidencias,
    titles_from_data=True
)

grafico_coincidencias.set_categories(
    categorias_temas
)

grafico_coincidencias.height = 10
grafico_coincidencias.width = 17
grafico_coincidencias.legend = None

ws_graficos.add_chart(
    grafico_coincidencias,
    "J4"
)


# ============================================================
# GRÁFICO 3: PÁGINAS CON COINCIDENCIAS POR TEMA
# ============================================================

grafico_paginas = BarChart()

grafico_paginas.type = "bar"
grafico_paginas.style = 10
grafico_paginas.title = "Páginas con coincidencias por tema"
grafico_paginas.y_axis.title = "Tema"
grafico_paginas.x_axis.title = "Número de páginas"

datos_paginas = Reference(
    ws_resumen,
    min_col=4,
    min_row=fila_temas_inicio,
    max_row=fila_temas_fin
)

grafico_paginas.add_data(
    datos_paginas,
    titles_from_data=True
)

grafico_paginas.set_categories(
    categorias_temas
)

grafico_paginas.height = 10
grafico_paginas.width = 17
grafico_paginas.legend = None

ws_graficos.add_chart(
    grafico_paginas,
    "A25"
)


# ============================================================
# GRÁFICO 4: PRINCIPALES ADMINISTRADOS
# ============================================================

grafico_administrados = BarChart()

grafico_administrados.type = "bar"
grafico_administrados.style = 10
grafico_administrados.title = (
    "Principales administrados por número de resoluciones"
)
grafico_administrados.y_axis.title = "Administrado"
grafico_administrados.x_axis.title = "Número de resoluciones"

# Mostrar como máximo los 15 administrados con más resoluciones.
# Esto evita saturar el gráfico cuando existen muchos administrados.

ultima_fila_administrados = min(
    ws_administrados.max_row,
    16
)

if ultima_fila_administrados >= 2:

    datos_administrados = Reference(
        ws_administrados,
        min_col=2,
        min_row=1,
        max_row=ultima_fila_administrados
    )

    categorias_administrados = Reference(
        ws_administrados,
        min_col=1,
        min_row=2,
        max_row=ultima_fila_administrados
    )

    grafico_administrados.add_data(
        datos_administrados,
        titles_from_data=True
    )

    grafico_administrados.set_categories(
        categorias_administrados
    )

    grafico_administrados.height = 10
    grafico_administrados.width = 17
    grafico_administrados.legend = None

    ws_graficos.add_chart(
        grafico_administrados,
        "J25"
    )


# ============================================================
# ORGANIZACIÓN VISUAL DE LA HOJA
# ============================================================

ws_graficos.column_dimensions["A"].width = 18
ws_graficos.column_dimensions["B"].width = 18
ws_graficos.column_dimensions["C"].width = 18
ws_graficos.column_dimensions["D"].width = 18
ws_graficos.column_dimensions["E"].width = 18
ws_graficos.column_dimensions["F"].width = 18
ws_graficos.column_dimensions["G"].width = 18
ws_graficos.column_dimensions["H"].width = 18
ws_graficos.column_dimensions["I"].width = 4
ws_graficos.column_dimensions["J"].width = 18
ws_graficos.column_dimensions["K"].width = 18
ws_graficos.column_dimensions["L"].width = 18
ws_graficos.column_dimensions["M"].width = 18
ws_graficos.column_dimensions["N"].width = 18
ws_graficos.column_dimensions["O"].width = 18
ws_graficos.column_dimensions["P"].width = 18
ws_graficos.column_dimensions["Q"].width = 18


# ============================================================
# FORMATO GENERAL
# ============================================================

for hoja in wb.worksheets:
    hoja.sheet_view.showGridLines = True
    hoja.freeze_panes = hoja.freeze_panes or "A2"


# ============================================================
# GUARDAR ANÁLISIS
# ============================================================

wb.save(ARCHIVO_ANALISIS)


# ============================================================
# MENSAJE FINAL
# ============================================================

segundos = time.time() - inicio_proceso
minutos = segundos / 60

print()
print("=" * 80)
print("ANÁLISIS TEMÁTICO GENERADO")
print("=" * 80)
print()
print(f"Archivo resultados: {ARCHIVO_RESULTADOS}")
print(f"Archivo análisis:   {ARCHIVO_ANALISIS}")
print()
print("Hojas creadas en el análisis:")
print("  1. Resumen")
print("  2. Resoluciones")
print("  3. Detalle")
print("  4. Administrados")
print("  5. Coocurrencia")
print("  6. Gráficos")
print()
print(f"Tiempo total: {minutos:.1f} minutos")
print()
print("Los PDF fueron procesados únicamente en memoria.")
print("No se guardó ningún PDF en disco.")
print()
print("=" * 80)
print("FIN DEL PROCESO")
print("=" * 80)
